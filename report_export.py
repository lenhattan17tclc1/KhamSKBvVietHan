"""Xuất báo cáo tổng hợp xét nghiệm theo đối tượng ra Excel và Word."""
import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import COL_TOTAL, fmt_date

FONT = "Times New Roman"


def _period(rep):
    return f"Từ ngày {fmt_date(rep.date_from)} đến ngày {fmt_date(rep.date_to)}"


def _today_line(place):
    d = dt.date.today()
    return f"{place}, ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"


def _rows(rep):
    """Các dòng của bảng tổng hợp: [STT, Tên XN, số liệu..., Tổng]."""
    cols = rep.columns + [COL_TOTAL]
    out = []
    for i, (t, v) in enumerate(rep.table.items(), 1):
        out.append([i, t] + [v[c] for c in cols])
    out.append(["", "Số người được xét nghiệm"] + [rep.total_people[c] for c in cols])
    return out


NOTE = ("Ghi chú: Số liệu tính theo số người (mỗi người chỉ tính 1 lần cho mỗi xét nghiệm). "
        "Người thuộc nhiều đối tượng được tính ở mỗi cột đối tượng tương ứng, "
        "vì vậy cột “Tổng số người” có thể nhỏ hơn tổng các cột đối tượng.")


# ---------------------------------------------------------------- Excel

def export_excel(rep, path, info):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tong hop"
    cols = rep.columns + [COL_TOTAL]
    ncol = 2 + len(cols)
    last = get_column_letter(ncol)
    thin = Side(style="thin")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    f = lambda **k: Font(name=FONT, size=k.pop("size", 12), **k)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    half = max(2, ncol // 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=half)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=half)
    ws.merge_cells(start_row=1, start_column=half + 1, end_row=1, end_column=ncol)
    ws.merge_cells(start_row=2, start_column=half + 1, end_row=2, end_column=ncol)
    ws.cell(1, 1, info["co_quan"].upper()).font = f()
    ws.cell(2, 1, info["don_vi"].upper()).font = f(bold=True)
    ws.cell(1, half + 1, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM").font = f(bold=True)
    ws.cell(2, half + 1, "Độc lập - Tự do - Hạnh phúc").font = f(bold=True, underline="single")
    for r in (1, 2):
        for c in (1, half + 1):
            ws.cell(r, c).alignment = center

    for r, text, font in ((4, info["tieu_de"].upper(), f(bold=True, size=14)),
                          (5, info["dia_ban"], f(italic=True)),
                          (6, _period(rep), f(italic=True))):
        ws.merge_cells(f"A{r}:{last}{r}")
        c = ws.cell(r, 1, text)
        c.font, c.alignment = font, center

    hr = 8
    head = ["STT", "Tên xét nghiệm"] + cols
    fill = PatternFill("solid", fgColor="DDEBF7")
    for j, h in enumerate(head, 1):
        c = ws.cell(hr, j, h)
        c.font, c.alignment, c.border, c.fill = f(bold=True, size=11), center, bd, fill
    ws.row_dimensions[hr].height = 48
    rows = _rows(rep)
    for i, vals in enumerate(rows, hr + 1):
        is_total = i == hr + len(rows)
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.border = bd
            c.font = f(bold=is_total or j == ncol, size=11)
            c.alignment = Alignment(horizontal="left" if j == 2 else "center", vertical="center", wrap_text=True)
            if is_total:
                c.fill = PatternFill("solid", fgColor="F2F2F2")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    for j in range(3, ncol + 1):
        ws.column_dimensions[get_column_letter(j)].width = 13

    r = hr + len(rows) + 2
    ws.merge_cells(f"A{r}:{last}{r}")
    c = ws.cell(r, 1, NOTE)
    c.font, c.alignment = f(italic=True, size=11), Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32

    r += 2
    ws.merge_cells(start_row=r, start_column=half + 1, end_row=r, end_column=ncol)
    c = ws.cell(r, half + 1, _today_line(info["noi_ky"]))
    c.font, c.alignment = f(italic=True), center
    r += 1
    for col1, col2, text in ((1, half, "NGƯỜI LẬP BIỂU"), (half + 1, ncol, "THỦ TRƯỞNG ĐƠN VỊ")):
        ws.merge_cells(start_row=r, start_column=col1, end_row=r, end_column=col2)
        c = ws.cell(r, col1, text)
        c.font, c.alignment = f(bold=True), center

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Sheet chi tiết từng người
    d = wb.create_sheet("Chi tiet")
    dh = ["STT", "Ngày lấy mẫu", "Số phiếu", "Họ tên", "Năm sinh", "Đối tượng", "Mức ghép DS", "Xét nghiệm đã làm"]
    for j, h in enumerate(dh, 1):
        c = d.cell(1, j, h)
        c.font, c.alignment, c.border, c.fill = f(bold=True, size=11), center, bd, fill
    for i, vals in enumerate(rep.details, 2):
        for j, v in enumerate([i - 1] + vals, 1):
            c = d.cell(i, j, v)
            c.font, c.border = f(size=11), bd
    for j, w in enumerate([6, 14, 10, 28, 10, 32, 15, 60], 1):
        d.column_dimensions[get_column_letter(j)].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = f"A1:H{len(rep.details) + 1}"
    wb.save(path)


# ---------------------------------------------------------------- Word

def export_word(rep, path, info):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.5))
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)

    def para(cell_or_doc, text, bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.CENTER, underline=False):
        p = cell_or_doc.add_paragraph()
        p.alignment = align
        r = p.add_run(text)
        r.bold, r.italic, r.underline = bold, italic, underline
        if size:
            r.font.size = Pt(size)
        return p

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    # Quốc hiệu
    t = doc.add_table(rows=1, cols=2)
    t.autofit = True
    left, right = t.rows[0].cells
    left._tc.remove(left.paragraphs[0]._p)
    right._tc.remove(right.paragraphs[0]._p)
    para(left, info["co_quan"].upper(), size=12)
    para(left, info["don_vi"].upper(), bold=True, size=12)
    para(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, size=12)
    para(right, "Độc lập - Tự do - Hạnh phúc", bold=True, underline=True, size=13)

    doc.add_paragraph()
    para(doc, info["tieu_de"].upper(), bold=True, size=15)
    para(doc, info["dia_ban"], italic=True)
    para(doc, _period(rep), italic=True)
    doc.add_paragraph()

    cols = rep.columns + [COL_TOTAL]
    head = ["STT", "Tên xét nghiệm"] + cols
    rows = _rows(rep)
    tb = doc.add_table(rows=1 + len(rows), cols=len(head))
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.1), Cm(5.6)] + [Cm((26.7 - 6.7) / len(cols))] * len(cols)
    for i, vals in enumerate([head] + rows):
        is_head, is_total = i == 0, i == len(rows)
        for j, v in enumerate(vals):
            c = tb.cell(i, j)
            c.width = widths[j]
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 1 and not is_head) else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(v))
            r.font.size = Pt(10.5 if is_head else 11.5)
            r.bold = is_head or is_total or j == len(vals) - 1
            if is_head:
                shade(c, "DDEBF7")
            elif is_total:
                shade(c, "F2F2F2")

    p = para(doc, NOTE, italic=True, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.space_before = Pt(6)

    doc.add_paragraph()
    sig = doc.add_table(rows=1, cols=2)
    a, b = sig.rows[0].cells
    a._tc.remove(a.paragraphs[0]._p)
    b._tc.remove(b.paragraphs[0]._p)
    para(a, " ")
    para(a, "NGƯỜI LẬP BIỂU", bold=True)
    para(a, "(Ký, ghi rõ họ tên)", italic=True, size=12)
    para(b, _today_line(info["noi_ky"]), italic=True)
    para(b, "THỦ TRƯỞNG ĐƠN VỊ", bold=True)
    para(b, "(Ký tên, đóng dấu)", italic=True, size=12)
    doc.save(path)
